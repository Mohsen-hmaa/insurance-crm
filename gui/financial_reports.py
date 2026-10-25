import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from gui.styles import (
    COLORS,
    configure_report_styles,
    finalize_report_window,
)
from gui.window_registry import SingleWindowMixin, register_window
from services.date_service import (
    convert_date_to_shamsi,
    convert_shamsi_to_date,
)
from services.financial_report_service import (
    get_collection_financial_report,
    get_total_payments_report,
)


class FinancialReports(SingleWindowMixin, tk.Toplevel):

    def __init__(self, parent, show_menu=True):
        super().__init__(parent)

        configure_report_styles()

        self.configure(bg=COLORS["surface"])
        self.title("گزارش‌های مالی")
        self.geometry("600x500")
        self.resizable(False, False)

        if show_menu:
            self.create_widgets()

    # =========================================================
    # Main UI
    # =========================================================
    def create_widgets(self):
        ttk.Label(
            self,
            text="گزارش‌های مالی",
            font=("Tahoma", 16, "bold"),
        ).pack(pady=20)

        financial_frame = ttk.LabelFrame(self, text="")
        financial_frame.pack(fill="x", padx=30, pady=10)

        buttons = [
            ("گزارش پرداختی ها از صندوق و پیش دریافتی", self.open_collection_report),
            ("گزارش کل پرداختی‌ها", self.open_total_payments_report),
            ("گزارش بدهکاران", self.open_debt_report),
            ("گزارش مانده حساب", self.open_balance_report),
            ("گزارش مالی یک بیمه‌نامه", self.open_policy_financial_report),
        ]
        for text, command in buttons:
            ttk.Button(
                financial_frame,
                text=text,
                command=command,
            ).pack(fill="x", padx=15, pady=5, ipady=10)

        finalize_report_window(self)

    # =========================================================
    # پنجره‌ی مشترک «بازه تاریخ وصول»
    # =========================================================
    def _open_date_range_window(
        self,
        window_title,
        heading,
        on_generate,
        key,
    ):
        """
        on_generate(from_date, to_date, window) با زدن «تهیه گزارش» صدا زده می‌شود.
        key: کلید یکتای گزارش؛ اگر پنجره‌اش باز باشد همان جلو می‌آید.
        """
        if self.reuse_window(key):
            return

        window = tk.Toplevel(self)
        register_window(key, window)
        window.configure(bg=COLORS["surface"])
        window.title(window_title)
        window.geometry("500x380")
        window.resizable(False, False)

        ttk.Label(
            window,
            text=heading,
            style="CommissionTitle.TLabel",
            anchor="e",
        ).pack(fill="x", padx=30, pady=20)

        filter_frame = ttk.LabelFrame(
            window,
            text="بازه تاریخ وصول",
            labelanchor="ne",
        )
        filter_frame.pack(fill="x", padx=30, pady=10)
        filter_frame.columnconfigure(0, weight=1)

        entries = []
        for row, label_text in enumerate(
            ("از تاریخ وصول:", "تا تاریخ وصول:")
        ):
            ttk.Label(
                filter_frame,
                text=label_text,
            ).grid(row=row, column=1, padx=10, pady=10, sticky="e")

            entry = ttk.Entry(filter_frame, width=20, justify="center")
            entry.grid(row=row, column=0, padx=10, pady=10, sticky="e")
            entries.append(entry)

        from_date_entry, to_date_entry = entries

        ttk.Button(
            window,
            text="تهیه گزارش",
            command=lambda: on_generate(
                from_date_entry.get().strip(),
                to_date_entry.get().strip(),
                window,
            ),
        ).pack(pady=20, ipadx=30, ipady=4)

        finalize_report_window(window)

    # =========================================================
    # گزارش پرداختی‌ها از صندوق و پیش‌دریافتی
    # =========================================================
    def open_collection_report(self):
        self._open_date_range_window(
            "گزارش پرداختی ها از صندوق و پیش دریافتی",
            "پرداختی‌ها از صندوق و پیش‌دریافتی",
            self.generate_collection_report,
            "financial_collection",
        )

    def generate_collection_report(self, from_date, to_date, window):
        self._generate_report(
            from_date,
            to_date,
            window,
            fetch=get_collection_financial_report,
            report_name="گزارش وصولی",
            sheet_title="گزارش وصولی",
            default_filename=f"حساب کتاب تا {to_date.replace('/', '')}.xlsx",
            empty_message="در بازه انتخاب‌شده وصولی موردی پیدا نشد.",
            success_message="فایل گزارش وصولی با موفقیت ذخیره شد.",
            with_summary=False,
        )

    # =========================================================
    # گزارش کل پرداختی‌ها
    # =========================================================
    def open_total_payments_report(self):
        self._open_date_range_window(
            "گزارش کل پرداختی‌ها",
            "کل پرداختی‌ها",
            self.generate_total_payments_report,
            "financial_total_payments",
        )

    def generate_total_payments_report(self, from_date, to_date, window):
        self._generate_report(
            from_date,
            to_date,
            window,
            fetch=get_total_payments_report,
            report_name="گزارش کل پرداختی‌ها",
            sheet_title="کل پرداختی‌ها",
            default_filename=(
                f"کل پرداختی ها تا {to_date.replace('/', '')}.xlsx"
            ),
            empty_message="در بازه انتخاب‌شده پرداختی‌ای پیدا نشد.",
            success_message="فایل گزارش کل پرداختی‌ها با موفقیت ذخیره شد.",
            with_summary=True,
        )

    # =========================================================
    # جریان مشترک: اعتبارسنجی، دریافت داده، ذخیره‌ی اکسل
    # =========================================================
    def _generate_report(
        self,
        from_date,
        to_date,
        window,
        fetch,
        report_name,
        sheet_title,
        default_filename,
        empty_message,
        success_message,
        with_summary,
    ):
        # اعتبارسنجی تاریخ‌ها
        # (سرویس برای ورودی نامعتبر فقط لیست خالی برمی‌گرداند)
        if not from_date or not to_date:
            messagebox.showwarning(
                "خطا",
                "لطفاً تاریخ شروع و پایان را وارد کنید.",
                parent=window,
            )
            return

        try:
            start = convert_shamsi_to_date(from_date)
            end = convert_shamsi_to_date(to_date)
        except Exception:
            start = end = None

        if not start or not end:
            messagebox.showerror(
                "خطا",
                "فرمت تاریخ وارد شده صحیح نیست.",
                parent=window,
            )
            return

        if start > end:
            messagebox.showwarning(
                "خطا",
                "تاریخ شروع نمی‌تواند بعد از تاریخ پایان باشد.",
                parent=window,
            )
            return

        try:
            collections = fetch(from_date, to_date)

            if not collections:
                messagebox.showinfo(
                    report_name,
                    empty_message,
                    parent=window,
                )
                return

            file_path = filedialog.asksaveasfilename(
                parent=window,
                title=f"ذخیره {report_name}",
                defaultextension=".xlsx",
                initialfile=default_filename,
                filetypes=[("Excel Files", "*.xlsx")],
            )
            if not file_path:
                return

            self._write_collection_excel(
                collections,
                file_path,
                sheet_title,
                with_summary,
            )

            messagebox.showinfo(
                report_name,
                success_message,
                parent=window,
            )

        except Exception as e:
            messagebox.showerror(
                "خطا",
                f"خطا در تهیه گزارش:\n{e}",
                parent=window,
            )

    # =========================================================
    # Excel
    # =========================================================
    @staticmethod
    def _write_collection_excel(
        collections,
        file_path,
        sheet_title,
        with_summary=False,
    ):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_title

        # راست به چپ
        sheet.sheet_view.rightToLeft = True

        # Styles
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        header_fill = PatternFill(fill_type="solid", fgColor="7370C7")
        even_fill = PatternFill(fill_type="solid", fgColor="E8EEF7")
        odd_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")

        header_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
        body_font = Font(name="Tahoma", size=10)
        bold_font = Font(name="Tahoma", size=10, bold=True)

        center_alignment = Alignment(horizontal="center", vertical="center")
        right_alignment = Alignment(horizontal="right", vertical="center")

        # Headers
        headers = [
            "ردیف",
            "شماره بیمه",
            "بازاریاب / معرف",
            "رشته",
            "بیمه‌گذار",
            "وصولی",
            "تاریخ وصول",
            "نوع دریافتی",
        ]
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
            cell.border = thin_border

        # Data
        for row_number, collection in enumerate(collections, start=2):
            policy = collection.policy

            # بازاریاب / معرف
            if policy and policy.marketer:
                marketer_or_referrer = policy.marketer.name
            elif policy and policy.referrer:
                marketer_or_referrer = policy.referrer
            else:
                marketer_or_referrer = "-"

            # بیمه‌گذار
            customer_name = "-"
            if policy and policy.customer:
                customer = policy.customer
                customer_name = (
                    f"{customer.name or ''} {customer.family or ''}"
                ).strip() or "-"

            # رشته
            insurance_type = "-"
            if policy and policy.insurance_type:
                insurance_type = (
                    policy.insurance_type.display_name or ""
                ).strip() or "-"

            values = [
                row_number - 1,
                policy.policy_number if policy else "-",
                marketer_or_referrer,
                insurance_type,
                customer_name,
                collection.settlement_amount or 0,
                convert_date_to_shamsi(collection.collection_date),
                collection.receipt_type or "-",
            ]

            row_fill = even_fill if (row_number - 1) % 2 == 0 else odd_fill

            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_number, column=column, value=value)
                cell.font = body_font
                cell.border = thin_border
                cell.fill = row_fill

                # رشته و بیمه‌گذار راست‌چین
                cell.alignment = (
                    right_alignment if column in (4, 5) else center_alignment
                )

                # فرمت عدد وصولی
                if column == 6:
                    cell.number_format = "#,##0"

        # Column widths
        column_widths = {
            "A": 8,
            "B": 18,
            "C": 25,
            "D": 25,
            "E": 28,
            "F": 18,
            "G": 15,
            "H": 18,
        }
        for column, width in column_widths.items():
            sheet.column_dimensions[column].width = width

        # Row heights
        sheet.row_dimensions[1].height = 25
        for row_number in range(2, len(collections) + 2):
            sheet.row_dimensions[row_number].height = 22

        # Freeze header
        sheet.freeze_panes = "A2"

        # ---------------------------------------------------------
        # شیت خلاصه: جمع به تفکیک نوع دریافتی
        # ---------------------------------------------------------
        if with_summary:
            summary = {}
            for collection in collections:
                receipt_type = collection.receipt_type or "-"
                item = summary.setdefault(receipt_type, [0, 0])
                item[0] += 1
                item[1] += collection.settlement_amount or 0

            summary_sheet = workbook.create_sheet("خلاصه")
            summary_sheet.sheet_view.rightToLeft = True

            for column, header in enumerate(
                ["نوع دریافتی", "تعداد", "جمع مبلغ"], start=1
            ):
                cell = summary_sheet.cell(row=1, column=column, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_alignment
                cell.border = thin_border

            rows = sorted(
                summary.items(),
                key=lambda pair: pair[1][1],
                reverse=True,
            )

            for row_number, (receipt_type, (count, amount)) in enumerate(
                rows, start=2
            ):
                row_fill = even_fill if row_number % 2 == 1 else odd_fill
                for column, value in enumerate(
                    [receipt_type, count, amount], start=1
                ):
                    cell = summary_sheet.cell(
                        row=row_number, column=column, value=value
                    )
                    cell.font = body_font
                    cell.border = thin_border
                    cell.fill = row_fill
                    cell.alignment = center_alignment
                    if column == 3:
                        cell.number_format = "#,##0"

            total_row = len(rows) + 2
            totals = [
                "جمع کل",
                sum(item[0] for item in summary.values()),
                sum(item[1] for item in summary.values()),
            ]
            for column, value in enumerate(totals, start=1):
                cell = summary_sheet.cell(
                    row=total_row, column=column, value=value
                )
                cell.font = bold_font
                cell.border = thin_border
                cell.fill = even_fill
                cell.alignment = center_alignment
                if column == 3:
                    cell.number_format = "#,##0"

            for column, width in {"A": 25, "B": 12, "C": 20}.items():
                summary_sheet.column_dimensions[column].width = width
            summary_sheet.row_dimensions[1].height = 25

        workbook.save(file_path)

    # =========================================================
    # Placeholder Reports
    # =========================================================
    def _not_connected(self, name):
        messagebox.showinfo(
            "گزارش مالی",
            f"گزارش {name} هنوز متصل نشده است.",
            parent=self,
        )

    def open_debt_report(self):
        self._not_connected("بدهکاران")

    def open_balance_report(self):
        self._not_connected("مانده حساب")

    def open_policy_financial_report(self):
        self._not_connected("مالی بیمه‌نامه")
