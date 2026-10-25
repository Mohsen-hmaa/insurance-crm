import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from services.insurance_type_service import load_root_insurance_types
from gui.insurance_selector import InsuranceSelector
from services.policy_service import (
    load_policies,
    delete_policy,
    get_policy_by_id
)
from gui.styles import COLORS, schedule_title_bar_color
from services.marketer_service import load_marketers
from gui.policy_form import open_policy_form
from gui.policy_details import open_policy_details
from gui.customer_selector import open_customer_selector
from services.date_service import convert_date_to_shamsi
from services.system_config_service import get_current_agent
from services.policy_filter_service import (
    filter_policies,
    DATE_CONDITIONS as FILTER_DATE_CONDITIONS,
    NUMBER_CONDITIONS as FILTER_NUMBER_CONDITIONS,
)
class PolicyList(tk.Frame):
    EDIT_ROLES = ("کاربر ارشد", "مدیر", "مدیر سیستم")
    # ورودی‌های متنی بعد از این مدت توقف تایپ جستجو می‌کنند (میلی‌ثانیه).
    SEARCH_DELAY_MS = 400
    # کلیدهایی که متن را عوض نمی‌کنند و جستجو لازم ندارند
    IGNORED_KEYS = frozenset({
        "Left", "Right", "Up", "Down", "Home", "End", "Prior", "Next",
        "Shift_L", "Shift_R", "Control_L", "Control_R",
        "Alt_L", "Alt_R", "Tab", "Caps_Lock", "Escape",
    })
    # فهرست شرط‌ها در سرویس فیلتر تعریف شده تا املای آن‌ها با منطق جستجو
    # همیشه یکی باشد (تاریخ‌ها: شامل «مابین»، حق‌بیمه و سهم صندوق: عددی)
    DATE_CONDITIONS = FILTER_DATE_CONDITIONS
    NUMBER_CONDITIONS = FILTER_NUMBER_CONDITIONS
    # شماره بیمه‌نامه فقط با شرط «برابر با» جستجو می‌شود
    POLICY_NUMBER_CONDITION = "برابر با"
    # مقدار «تریستیت» رادیوبوتن‌ها. پیش‌فرض Tk رشته‌ی خالی است و وقتی متغیر
    # خالی باشد هر دو دکمه پُر (انتخاب‌شده) نمایش داده می‌شوند.
    RADIO_NONE = "__none__"
    # نام فیلتر -> (کمبوی شرط، ورودی اول، ورودی دوم برای «مابین»)
    DATE_FILTERS = {
        "issue_date": (
            "issue_condition_combo",
            "issue_date_search_entry",
            "issue_date_search_entry2",
        ),
        "start_date": (
            "start_condition_combo",
            "start_date_entry",
            "start_date_entry2",
        ),
        "end_date": (
            "end_condition_combo",
            "end_date_entry",
            "end_date_entry2",
        ),
    }
    def __init__(self, parent, role="کاربر"):
        super().__init__(
            parent,
            bg=COLORS["background"]
        )
        self.parent = parent
        self.role = role
        self.selected_marketer = None
        self.selected_customers = []
        self.selected_insurance_type = None
        self.main_insurance_type = None
        self.search_window = None
        self.filtered_policies = None
        self.all_policies = []
        self.marketers = []
        self.search_vars = {}
        self.search_frames = {}
        self._search_job = None
        self._row_cache = {}
        # (شناسه بیمه‌نامه، مقادیر ردیف) به همان ترتیبی که در جدول هست
        self._displayed_rows = []
        self.current_agent = get_current_agent()
        self.bind("<Destroy>", self._on_destroy, add="+")
        self.create_widgets()
        self.load_data()
    # ==================================================
    # جستجوی تأخیری (Debounce)
    # ==================================================
    def schedule_policy_search(self, event=None):
        """
        جستجو را به بعد از توقف تایپ موکول می‌کند تا با هر کاراکتر
        یک جستجوی کامل روی همه‌ی رکوردها اجرا نشود.
        """
        if (
            event is not None
            and getattr(event, "keysym", None) in self.IGNORED_KEYS
        ):
            return
        self._cancel_pending_search()
        self._search_job = self.after(
            self.SEARCH_DELAY_MS,
            self._run_scheduled_search,
        )
    def _run_scheduled_search(self):
        self._search_job = None
        self.refresh_policy_search_results()
    def _cancel_pending_search(self):
        if self._search_job is not None:
            try:
                self.after_cancel(self._search_job)
            except tk.TclError:
                pass
            self._search_job = None
    def _on_destroy(self, event):
        if event.widget is self:
            self._cancel_pending_search()
    # ==================================================
    # ابزارهای کمکی
    # ==================================================
    @staticmethod
    def _set_readonly_text(entry, text=""):
        entry.config(state="normal")
        entry.delete(0, tk.END)
        if text:
            entry.insert(0, text)
        entry.config(state="readonly")
    def _date_widgets(self, filter_name):
        combo_attr, entry_attr, entry2_attr = self.DATE_FILTERS[filter_name]
        return (
            getattr(self, combo_attr),
            getattr(self, entry_attr),
            getattr(self, entry2_attr),
        )
    @staticmethod
    def _customer_text(customer):
        return (
            f"{customer.name} "
            f"{customer.family or ''} "
            f"{customer.customer_code or ''}"
        ).strip()
    def _selected_policy_id(self):
        selected = self.tree.selection()
        if not selected:
            return None
        try:
            return int(selected[0])
        except (ValueError, TypeError):
            return None
    # ==================================================
    # ساخت صفحه
    # ==================================================
    def create_widgets(self):
        # ==================================================
        # Header
        # ==================================================
        header_frame = tk.Frame(
            self,
            bg=COLORS["background"]
        )
        header_frame.pack(fill="x", pady=(0, 15))
        tk.Label(
            header_frame,
            text="لیست بیمه‌نامه‌ها",
            bg=COLORS["background"],
            fg=COLORS["text"],
            font=("Tahoma", 18, "bold"),
            anchor="e"
        ).pack(side="right")
        # ==================================================
        # Columns
        # ==================================================
        columns = (
            "mortgage",
            "premium",
            "marketer",
            "referrer",
            "end_date",
            "start_date",
            "issue_date",
            "insurance_type",
            "endorsement_number",
            "policy_number",
            "policy_type",
            "customer",
            "row",
        )
        widths = {
            "mortgage": 50,
            "premium": 90,
            "marketer": 70,
            "referrer": 240,
            "end_date": 80,
            "start_date": 80,
            "issue_date": 80,
            "insurance_type": 240,
            "endorsement_number": 50,
            "policy_number": 60,
            "policy_type": 60,
            "customer": 240,
            "row": 50,
        }
        # ==================================================
        # Treeview Style
        # ==================================================
        style = ttk.Style()
        style.map(
            "Treeview",
            background=[("selected", "#3498DB")],
            foreground=[("selected", "white")]
        )
        # ==================================================
        # Table Frame
        # ==================================================
        table_frame = tk.Frame(
            self,
            bg=COLORS["background"]
        )
        table_frame.pack(fill="both", expand=True, padx=10, pady=10)
        # ==================================================
        # Treeview
        # ==================================================
        self.tree = ttk.Treeview(
            table_frame,
            columns=columns,
            show="headings"
        )
        # Row Colors
        self.tree.tag_configure("even", background="#E0E0E0")
        self.tree.tag_configure("odd", background="#FFFFFF")
        self.tree.tag_configure(
            "endorsement_even",
            background="#E0E0E0",
            foreground="#C91F28"
        )
        self.tree.tag_configure(
            "endorsement_odd",
            background="#FFFFFF",
            foreground="#C91F28"
        )
        # Headers
        headers = {
            "mortgage": "مرهوناتی",
            "premium": "حق‌بیمه",
            "marketer": "بازاریاب",
            "referrer": "معرف",
            "end_date": "تاریخ پایان",
            "start_date": "تاریخ شروع",
            "issue_date": "تاریخ صدور",
            "insurance_type": "رشته بیمه",
            "endorsement_number": "الحاقیه",
            "policy_number": "ش بیمه",
            "policy_type": "نوع",
            "customer": "بیمه‌گذار",
            "row": "ردیف",
        }
        for column, text in headers.items():
            self.tree.heading(column, text=text, anchor="e")
            self.tree.column(
                column,
                width=widths[column],
                minwidth=widths[column],
                anchor="e",
                stretch=False
            )
        # Scrollbars
        y_scrollbar = ttk.Scrollbar(
            table_frame,
            orient="vertical",
            command=self.tree.yview
        )
        x_scrollbar = ttk.Scrollbar(
            table_frame,
            orient="horizontal",
            command=self.tree.xview
        )
        self.tree.configure(
            yscrollcommand=y_scrollbar.set,
            xscrollcommand=x_scrollbar.set
        )
        # Layout
        self.tree.grid(row=0, column=0, sticky="nsew")
        y_scrollbar.grid(row=0, column=1, sticky="ns")
        x_scrollbar.grid(row=1, column=0, sticky="ew")
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)
        # ==================================================
        # Buttons
        # ==================================================
        button_frame = tk.Frame(self, bg=COLORS["background"])
        button_frame.pack(pady=10)
        self.edit_button = tk.Button(
            button_frame,
            text="ویرایش بیمه‌نامه",
            background="#CDCCCC",
            width=20,
            command=self.edit_policy
        )
        self.edit_button.pack(side="right", padx=20)
        self.delete_button = tk.Button(
            button_frame,
            text="حذف بیمه‌نامه",
            background="#CDCCCC",
            width=20,
            command=self.delete_selected_policy
        )
        self.delete_button.pack(side="right", padx=20)
        self.search_button = tk.Button(
            button_frame,
            text="جستجوی بیمه‌نامه",
            background="#CDCCCC",
            width=20,
            command=self.open_search_form
        )
        self.search_button.pack(side="right", padx=20)
        self.details_button = tk.Button(
            button_frame,
            text="جزئیات بیمه‌نامه",
            background="#CDCCCC",
            width=20,
            command=self.open_policy_details
        )
        self.details_button.pack(side="right", padx=20)
        self.export_excel_button = tk.Button(
            button_frame,
            text="انتقال به فایل اکسل",
            background="#CDCCCC",
            width=20,
            command=self.export_to_excel
        )
        self.export_excel_button.pack(side="right", padx=20)
        # تعیین دسترسی دکمه‌ها براساس نقش کاربر
        if self.role == "کاربر":
            self.edit_button.pack_forget()
        if self.role != "مدیر سیستم":
            self.delete_button.pack_forget()
    # ==================================================
    # داده‌ها
    # ==================================================
    def _reload_policies(self):
        """بارگذاری مجدد از دیتابیس؛ cache ردیف‌ها هم باطل می‌شود."""
        self.all_policies = load_policies()
        self._row_cache.clear()
    def _show_current(self):
        """نمایش وضعیت فعلی (با فیلتر فعال، اگر وجود داشته باشد)."""
        if self.filtered_policies is not None:
            self.search_policy(close_window=False)
        else:
            self.load_data()
    def load_data(self, policies=None):
        if policies is None:
            if not self.all_policies:
                self._reload_policies()
            policies = self.all_policies
        self._render_policies(policies)
    def refresh_data(self):
        self._reload_policies()
        self._show_current()
    def _build_row_values(self, policy):
        """
        مقادیر نمایشی هر ردیف (بدون شماره ردیف) فقط یک بار حساب و cache
        می‌شود. شماره ردیف به موقعیت در جدول بستگی دارد و جدا اضافه می‌شود.
        """
        cached = self._row_cache.get(policy.id)
        if cached is not None:
            return cached
        mortgage = "✓" if policy.is_mortgage else "-"
        referrer = policy.referrer
        if (
            referrer
            and referrer.strip() == (self.current_agent or "").strip()
        ):
            referrer_text = ""
        else:
            referrer_text = referrer if referrer else "------"
        customer_text = "------"
        if policy.customer:
            customer_text = self._customer_text(policy.customer)
        insurance_type_text = "------"
        if policy.insurance_type:
            insurance_type_text = policy.insurance_type.display_name.strip()
        def shamsi(value):
            return convert_date_to_shamsi(value) if value else "-----"
        values = (
            mortgage,
            f"{policy.premium or 0:,}",
            policy.marketer.name.strip() if policy.marketer else "-----",
            referrer_text,
            shamsi(policy.end_date),
            shamsi(policy.start_date),
            shamsi(policy.issue_date),
            insurance_type_text,
            policy.endorsement_number,
            policy.policy_number,
            "بیمه‌نامه" if policy.endorsement_number == 0 else "الحاقیه",
            customer_text,
        )
        self._row_cache[policy.id] = values
        return values
    def _render_policies(self, policies):
        self.tree.delete(*self.tree.get_children())
        self._displayed_rows = []
        # بیمه‌نامه‌های بدون تاریخ صدور اول می‌آیند. مقایسه‌ی مستقیم
        # None با تاریخ در پایتون خطا می‌دهد، پس جدا مرتب می‌شوند.
        undated = [p for p in policies if not p.issue_date]
        dated = sorted(
            (p for p in policies if p.issue_date),
            key=lambda policy: policy.issue_date
        )
        for index, policy in enumerate(undated + dated):
            even = index % 2 == 0
            if policy.endorsement_number != 0:
                tag = "endorsement_even" if even else "endorsement_odd"
            else:
                tag = "even" if even else "odd"
            values = self._build_row_values(policy) + (index + 1,)
            self._displayed_rows.append((policy.id, values))
            self.tree.insert(
                "",
                "end",
                iid=str(policy.id),
                values=values,
                tags=(tag,)
            )
        self.tree.yview_moveto(1)
        self.tree.xview_moveto(1)
    # ==================================================
    # ویرایش / جزئیات / حذف
    # ==================================================
    def edit_policy(self):
        if self.role not in self.EDIT_ROLES:
            return
        policy_id = self._selected_policy_id()
        if policy_id is None:
            messagebox.showwarning(
                "هشدار",
                "لطفاً یک بیمه‌نامه را انتخاب کنید",
                parent=self
            )
            return
        policy = get_policy_by_id(policy_id)
        if not policy:
            messagebox.showerror(
                "خطا",
                "بیمه‌نامه پیدا نشد",
                parent=self
            )
            return
        result = open_policy_form(policy, self)
        if result:
            # اطلاعات را مجدداً از دیتابیس بخوان و فیلتر فعال را دوباره اعمال کن
            self._reload_policies()
            self._show_current()
    def open_policy_details(self):
        policy_id = self._selected_policy_id()
        if policy_id is None:
            messagebox.showerror(
                "خطا",
                "لطفاً یک بیمه‌نامه را انتخاب کنید.",
                parent=self
            )
            return
        open_policy_details(self, policy_id)
    def delete_selected_policy(self):
        if self.role != "مدیر سیستم":
            return
        policy_id = self._selected_policy_id()
        if policy_id is None:
            messagebox.showerror(
                "خطا",
                "بیمه نامه ای انتخاب نشده",
                parent=self
            )
            return
        confirmation = messagebox.askyesno(
            "تایید حذف بیمه نامه",
            "آیا از حذف بیمه نامه اطمینان دارید؟",
            parent=self
        )
        if not confirmation:
            return
        try:
            delete_policy(policy_id)
        except Exception as e:
            messagebox.showerror(
                "خطا",
                f"حذف بیمه نامه انجام نشد:\n{e}",
                parent=self
            )
            return
        # چون جستجو از داده‌ی cache می‌خواند، بیمه‌نامه‌ی حذف‌شده باید از
        # همه‌ی فهرست‌ها هم برداشته شود، وگرنه در جستجوی بعدی برمی‌گردد.
        self.all_policies = [
            policy for policy in self.all_policies
            if policy.id != policy_id
        ]
        if self.filtered_policies is not None:
            self.filtered_policies = [
                policy for policy in self.filtered_policies
                if policy.id != policy_id
            ]
        self._row_cache.pop(policy_id, None)
        self._displayed_rows = [
            row for row in self._displayed_rows if row[0] != policy_id
        ]
        self.tree.delete(str(policy_id))
        messagebox.showinfo(
            "موفق",
            "بیمه نامه با موفقیت حذف گردید",
            parent=self
        )
    # ==================================================
    # اکسل
    # ==================================================
    def export_to_excel(self):
        # ردیف‌های فعلی جدول؛ شامل نتایج جستجو. مقادیر از خود داده خوانده
        # می‌شود (نه از Treeview) تا مثلاً صفرهای ابتدای شماره بیمه حفظ شود.
        rows = [values for _, values in self._displayed_rows]
        if not rows:
            messagebox.showwarning(
                "هشدار",
                "ردیفی برای انتقال به اکسل وجود ندارد.",
                parent=self
            )
            return
        file_path = filedialog.asksaveasfilename(
            parent=self,
            title="ذخیره فایل اکسل بیمه‌نامه‌ها",
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            initialfile="لیست_بیمه‌نامه‌ها.xlsx"
        )
        if not file_path:
            return
        try:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "بیمه‌نامه‌ها"
            sheet.sheet_view.rightToLeft = True
            # ستون‌ها برای شیت راست‌به‌چپ معکوس می‌شوند
            columns = list(self.tree["columns"])[::-1]
            headers = [
                self.tree.heading(column, "text")
                for column in columns
            ]
            sheet.append(headers)
            for values in rows:
                sheet.append(list(values)[::-1])
            # قالب‌بندی
            header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )
            header_font = Font(name="Tahoma", bold=True)
            body_font = Font(name="Tahoma", size=10)
            header_alignment = Alignment(horizontal="center", vertical="center")
            body_alignment = Alignment(horizontal="right", vertical="center")
            for cell in sheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = body_font
                    cell.alignment = body_alignment
                    cell.border = thin_border
            for column_cells in sheet.columns:
                column_letter = column_cells[0].column_letter
                max_length = max(
                    len(str(cell.value or ""))
                    for cell in column_cells
                )
                sheet.column_dimensions[column_letter].width = min(
                    max(max_length + 3, 12),
                    35
                )
            sheet.freeze_panes = "A2"
            workbook.save(file_path)
            messagebox.showinfo(
                "موفق",
                f"فایل اکسل با موفقیت ذخیره شد.\n"
                f"تعداد ردیف‌ها: {len(rows)}",
                parent=self
            )
        except Exception as e:
            messagebox.showerror(
                "خطا",
                f"انتقال اطلاعات به اکسل انجام نشد:\n{e}",
                parent=self
            )
    # ==================================================
    # جستجو
    # ==================================================
    def search_policy(self, close_window=True):
        # از داده‌ی cache جستجو می‌شود؛ بارگذاری مجدد از دیتابیس فقط بعد
        # از ویرایش یا با refresh_data انجام می‌شود.
        policies = self.all_policies
        customer_ids = None
        if self.search_vars["customer"].get():
            if self.selected_customers:
                customer_ids = [
                    customer.id
                    for customer in self.selected_customers
                ]
        policy_number = ""
        if self.search_vars["policy_number"].get():
            policy_number = self.policy_number_entry.get().strip()
        insurance_type_id = None
        if self.search_vars["insurance_type"].get():
            if self.selected_insurance_type:
                insurance_type_id = self.selected_insurance_type.id
        date_filters = {}
        for name in self.DATE_FILTERS:
            date_filters[name] = None
            if self.search_vars[name].get():
                combo, entry, entry2 = self._date_widgets(name)
                date_filters[name] = (combo.get(), entry.get(), entry2.get())
        referrer_condition = None
        referrer = ""
        if self.search_vars["referrer"].get():
            referrer_condition = self.referrer_condition_combo.get()
            referrer = self.referrer_entry.get().strip()
        marketer_condition = None
        marketer_id = None
        if self.search_vars["marketer"].get():
            marketer_condition = self.marketer_condition_combo.get()
            if self.selected_marketer:
                marketer_id = self.selected_marketer.id
        record_type = None
        if self.search_vars["record_type"].get():
            record_type = self.record_type_search_var.get() or None
        premium = None
        premium_condition = "برابر با"
        if self.search_vars["premium"].get():
            premium_text = (
                self.premium_entry.get().replace(",", "").strip()
            )
            premium_condition = self.premium_condition_combo.get()
            if premium_text:
                try:
                    premium = int(premium_text)
                except ValueError:
                    premium = None
        fund = None
        fund_condition = "برابر با"
        if self.search_vars["fund"].get():
            fund_text = (
                self.fund_search_entry.get().replace(",", "").strip()
            )
            fund_condition = self.fund_condition_combo.get()
            if fund_text:
                try:
                    fund = int(fund_text)
                except ValueError:
                    fund = None
        mortgage = None
        if self.search_vars["mortgage"].get():
            mortgage = self.mortgage_search_var.get() or None
        filtered_policies = filter_policies(
            policies,
            customer_ids=customer_ids,
            policy_number=policy_number,
            policy_number_condition=self.POLICY_NUMBER_CONDITION,
            insurance_type_id=insurance_type_id,
            issue_date=date_filters["issue_date"],
            start_date=date_filters["start_date"],
            end_date=date_filters["end_date"],
            referrer_condition=referrer_condition,
            referrer=referrer,
            marketer_condition=marketer_condition,
            marketer_id=marketer_id,
            record_type=record_type,
            premium=premium,
            premium_condition=premium_condition,
            fund=fund,
            fund_condition=fund_condition,
            mortgage=mortgage,
            current_agent=self.current_agent,
        )
        self.filtered_policies = filtered_policies
        if close_window:
            self.search_window.withdraw()
        self.load_data(filtered_policies)
    def refresh_policy_search_results(self, event=None):
        # اگر جستجوی تأخیری در انتظار بود، دیگر لازم نیست
        self._cancel_pending_search()
        # تاریخ باید یا خالی باشد یا کامل (YYYY/MM/DD)
        for name in self.DATE_FILTERS:
            if not self.search_vars[name].get():
                continue
            _, entry, entry2 = self._date_widgets(name)
            for date_entry in (entry, entry2):
                if len(date_entry.get().strip()) not in (0, 10):
                    return
        self.search_policy(close_window=False)
    # ==================================================
    # پنجره‌ی جستجو
    # ==================================================
    def open_search_form(self):
        if (
            self.search_window is not None
            and self.search_window.winfo_exists()
        ):
            self.search_window.deiconify()
            self.search_window.lift()
            return
        self.search_window = tk.Toplevel(self)
        self.search_window.protocol(
            "WM_DELETE_WINDOW",
            self.search_window.withdraw
        )
        self.search_window.title("جستجوی بیمه‌نامه")
        self.search_window.geometry("600x600")
        schedule_title_bar_color(self.search_window)
        tk.Label(
            self.search_window,
            text=":جستجو بر اساس",
            font=("Arial", 12, "bold")
        ).pack(pady=10, padx=10, anchor="e")
        # فریم اصلی برای راست‌چین کردن معیارها
        container = tk.Frame(self.search_window)
        container.pack(anchor="e", padx=20, pady=10)
        criteria_frame = tk.Frame(container)
        criteria_frame.pack()
        criteria = [
            ("customer", "بیمه‌گذار"),
            ("policy_number", "شماره بیمه‌نامه"),
            ("insurance_type", "رشته بیمه"),
            ("issue_date", "تاریخ صدور"),
            ("start_date", "تاریخ شروع"),
            ("end_date", "تاریخ پایان"),
            ("referrer", "معرف"),
            ("marketer", "بازاریاب"),
            ("record_type", "نوع رکورد"),
            ("premium", "حق‌بیمه"),
            ("fund", "سهم صندوق"),
            ("mortgage", "وضعیت مرهوناتی"),
        ]
        builders = {
            "customer": self.create_customer_search,
            "policy_number": self._create_policy_number_search,
            "insurance_type": self.create_insurance_type_search,
            "issue_date": lambda f: self._create_date_search(f, "issue_date"),
            "start_date": lambda f: self._create_date_search(f, "start_date"),
            "end_date": lambda f: self._create_date_search(f, "end_date"),
            "referrer": self._create_referrer_search,
            "marketer": self._create_marketer_search,
            "record_type": self._create_record_type_search,
            "premium": self._create_premium_search,
            "fund": self._create_fund_search,
            "mortgage": self._create_mortgage_search,
        }
        self.search_vars = {}
        self.search_frames = {}
        for row, (name, text) in enumerate(criteria):
            variable = tk.BooleanVar(
                master=self.search_window,
                value=False
            )
            self.search_vars[name] = variable
            # فریم فیلد جستجوی هر معیار
            field_frame = tk.Frame(criteria_frame)
            self.search_frames[name] = field_frame
            builders[name](field_frame)
            # چک‌باکس
            tk.Checkbutton(
                criteria_frame,
                variable=variable,
                command=lambda n=name: self.toggle_search_field(n)
            ).grid(row=row, column=1, padx=5, pady=4, sticky="e")
            # عنوان
            tk.Label(
                criteria_frame,
                text=text
            ).grid(row=row, column=2, padx=5, pady=4, sticky="e")
            # فریم جستجو؛ فعلاً مخفی
            field_frame.grid(row=row, column=0, padx=10, pady=4)
            field_frame.grid_remove()
        tk.Button(
            self.search_window,
            text="تایید",
            width=20,
            background="#CDCCCC",
            command=self.search_policy
        ).pack(pady=15)
        tk.Button(
            self.search_window,
            text="حذف همه فیلترها",
            width=20,
            background="#CDCCCC",
            command=self.clear_search_filters
        ).pack(pady=5)
    # ==================================================
    # سازنده‌ی هر فیلتر
    # ==================================================
    def _create_policy_number_search(self, frame):
        self.policy_number_entry = tk.Entry(frame, width=25)
        self.policy_number_entry.bind(
            "<KeyRelease>",
            self.schedule_policy_search
        )
        self.policy_number_entry.pack(side="right", padx=5)
    def _create_date_search(self, frame, filter_name):
        combo_attr, entry_attr, entry2_attr = self.DATE_FILTERS[filter_name]
        combo = ttk.Combobox(
            frame,
            width=15,
            values=self.DATE_CONDITIONS,
            state="readonly"
        )
        combo.current(0)
        combo.bind(
            "<<ComboboxSelected>>",
            lambda event, n=filter_name: self.change_date_condition(n)
        )
        combo.pack(side="right", padx=5)
        entry = tk.Entry(frame, width=15)
        entry.bind("<KeyRelease>", self.schedule_policy_search)
        entry.pack(side="right", padx=5)
        # ورودی دوم فقط برای شرط «مابین» نمایش داده می‌شود
        entry2 = tk.Entry(frame, width=15)
        entry2.bind("<KeyRelease>", self.schedule_policy_search)
        setattr(self, combo_attr, combo)
        setattr(self, entry_attr, entry)
        setattr(self, entry2_attr, entry2)
    def _create_referrer_search(self, frame):
        self.referrer_condition_combo = ttk.Combobox(
            frame,
            width=15,
            values=["معرف دارد", "بدون معرف", "برابر با", "شامل"],
            state="readonly"
        )
        self.referrer_condition_combo.current(3)
        self.referrer_condition_combo.bind(
            "<<ComboboxSelected>>",
            self.change_referrer_entry
        )
        self.referrer_condition_combo.pack(side="right", padx=5)
        self.referrer_entry = tk.Entry(frame, width=25, state="normal")
        self.referrer_entry.bind("<KeyRelease>", self.schedule_policy_search)
        self.referrer_entry.pack(side="right", padx=5)
    def _create_marketer_search(self, frame):
        self.marketer_condition_combo = ttk.Combobox(
            frame,
            width=15,
            values=["بازاریاب دارد", "بدون بازاریاب", "برابر با"],
            state="readonly"
        )
        self.marketer_condition_combo.current(2)
        self.marketer_condition_combo.bind(
            "<<ComboboxSelected>>",
            self.change_marketer_combo
        )
        self.marketer_condition_combo.pack(side="right", padx=5)
        # readonly: فقط از فهرست قابل انتخاب باشد، نه تایپ آزاد
        self.marketer_combo = ttk.Combobox(
            frame,
            width=25,
            state="readonly"
        )
        self.marketer_combo.bind(
            "<<ComboboxSelected>>",
            self.select_marketer
        )
        self.marketer_combo.pack(side="right", padx=5)
        # فهرست بازاریاب‌ها فقط یک بار از دیتابیس خوانده می‌شود
        self.marketers = load_marketers()
        self.marketer_combo["values"] = [
            marketer.name for marketer in self.marketers
        ]
    def _create_record_type_search(self, frame):
        self.record_type_search_var = tk.StringVar(
            master=self.search_window
        )
        tk.Radiobutton(
            frame,
            tristatevalue=self.RADIO_NONE,
            text="بیمه‌نامه",
            variable=self.record_type_search_var,
            value="بیمه‌نامه",
            command=self.refresh_policy_search_results
        ).pack(side="right")
        tk.Radiobutton(
            frame,
            tristatevalue=self.RADIO_NONE,
            text="الحاقیه",
            variable=self.record_type_search_var,
            value="الحاقیه",
            command=self.refresh_policy_search_results
        ).pack(side="right", padx=5)
    def _create_premium_search(self, frame):
        self.premium_condition_combo = ttk.Combobox(
            frame,
            width=15,
            values=self.NUMBER_CONDITIONS,
            state="readonly"
        )
        self.premium_condition_combo.current(0)
        self.premium_condition_combo.bind(
            "<<ComboboxSelected>>",
            self.refresh_policy_search_results
        )
        self.premium_condition_combo.pack(side="right", padx=5)
        self.premium_entry = tk.Entry(frame, width=25)
        self.premium_entry.bind("<KeyRelease>", self.format_premium)
        self.premium_entry.pack(side="right", padx=5)
    def _create_fund_search(self, frame):
        self.fund_condition_combo = ttk.Combobox(
            frame,
            width=15,
            values=self.NUMBER_CONDITIONS,
            state="readonly"
        )
        self.fund_condition_combo.current(0)
        self.fund_condition_combo.bind(
            "<<ComboboxSelected>>",
            self.refresh_policy_search_results
        )
        self.fund_condition_combo.pack(side="right", padx=5)
        self.fund_search_entry = tk.Entry(frame, width=25)
        self.fund_search_entry.bind(
            "<KeyRelease>",
            self.schedule_policy_search
        )
        self.fund_search_entry.pack(side="right", padx=5)
    def _create_mortgage_search(self, frame):
        self.mortgage_search_var = tk.StringVar(
            master=self.search_window
        )
        tk.Radiobutton(
            frame,
            tristatevalue=self.RADIO_NONE,
            text="غیر مرهوناتی",
            variable=self.mortgage_search_var,
            value="غیر مرهوناتی",
            command=self.refresh_policy_search_results
        ).pack(side="right")
        tk.Radiobutton(
            frame,
            tristatevalue=self.RADIO_NONE,
            text="مرهوناتی",
            variable=self.mortgage_search_var,
            value="مرهوناتی",
            command=self.refresh_policy_search_results
        ).pack(side="right", padx=5)
    # ==================================================
    # بیمه‌گذار (انتخاب چندتایی)
    # ==================================================
    def create_customer_search(self, frame):
        tk.Label(
            frame,
            text=": بیمه‌گذار",
            anchor="e"
        ).pack(side="right", padx=5)
        self.customer_search_entry = tk.Entry(
            frame,
            width=25,
            state="readonly"
        )
        self.customer_search_entry.pack(side="right", padx=5)
        self.customer_search_button = tk.Button(
            frame,
            text="انتخاب",
            background="#CDCCCC",
            command=self.customer_selector
        )
        self.customer_search_button.pack(side="right", padx=5)
    def customer_selector(self):
        open_customer_selector(
            self.search_window,
            self.customer_selected,
            multi_select=True
        )
    def customer_selected(self, customers):
        self.selected_customers = customers or []
        self._set_readonly_text(
            self.customer_search_entry,
            "، ".join(
                self._customer_text(customer)
                for customer in self.selected_customers
            )
        )
        self.search_policy(close_window=False)
    # ==================================================
    # رشته بیمه
    # ==================================================
    def create_insurance_type_search(self, frame):
        tk.Label(
            frame,
            text=": نوع بیمه",
            anchor="e"
        ).pack(side="right", padx=5)
        self.insurance_type_search_entry = tk.Entry(
            frame,
            width=25,
            state="readonly"
        )
        self.insurance_type_search_entry.pack(side="right", padx=5)
        self.insurance_type_search_button = tk.Button(
            frame,
            text="انتخاب",
            background="#CDCCCC",
            command=self.open_insurance_selector
        )
        self.insurance_type_search_button.pack(side="right", padx=5)
    def open_insurance_selector(self):
        root_types = load_root_insurance_types()
        InsuranceSelector(
            self.search_window,
            root_types,
            self.insurance_selected
        )
    def insurance_selected(self, selected_path):
        if not selected_path:
            return
        self.selected_insurance_type = selected_path[-1]
        self.main_insurance_type = selected_path[0]
        self._set_readonly_text(
            self.insurance_type_search_entry,
            selected_path[-1].display_name
        )
        self.refresh_policy_search_results()
    # ==================================================
    # تغییر فیلترها
    # ==================================================
    def _filter_has_value(self, name):
        """آیا فیلتر فعال‌شده الان مقدار یا شرطی دارد که جستجو را عوض کند؟"""
        if name == "customer":
            return bool(self.selected_customers)
        if name == "policy_number":
            return bool(self.policy_number_entry.get().strip())
        if name == "insurance_type":
            return bool(self.selected_insurance_type)
        if name in self.DATE_FILTERS:
            _, entry, _ = self._date_widgets(name)
            return bool(entry.get().strip())
        if name == "referrer":
            # «معرف دارد» و «بدون معرف» مقدار متنی نمی‌خواهند
            if self.referrer_condition_combo.get() in (
                "معرف دارد",
                "بدون معرف"
            ):
                return True
            return bool(self.referrer_entry.get().strip())
        if name == "marketer":
            # «بازاریاب دارد» و «بدون بازاریاب» انتخاب بازاریاب نمی‌خواهند
            if self.marketer_condition_combo.get() in (
                "بازاریاب دارد",
                "بدون بازاریاب"
            ):
                return True
            return bool(self.selected_marketer)
        if name == "record_type":
            return bool(self.record_type_search_var.get().strip())
        if name == "premium":
            return bool(self.premium_entry.get().strip())
        if name == "fund":
            return bool(self.fund_search_entry.get().strip())
        if name == "mortgage":
            return bool(self.mortgage_search_var.get().strip())
        return False
    def toggle_search_field(self, name):
        frame = self.search_frames[name]
        variable = self.search_vars[name]
        if variable.get():
            frame.grid()
        else:
            frame.grid_remove()
        self._cancel_pending_search()
        # اگر فیلتر خاموش شد، نتیجه جستجو باید به‌روز شود
        if not variable.get():
            self.search_policy(close_window=False)
            return
        if self._filter_has_value(name):
            self.search_policy(close_window=False)
    def change_date_condition(self, filter_name):
        combo, _, entry2 = self._date_widgets(filter_name)
        if combo.get() == "مابین":
            entry2.pack(side="right", padx=5)
        else:
            entry2.delete(0, tk.END)
            entry2.pack_forget()
        self.refresh_policy_search_results()
    def change_referrer_entry(self, event=None):
        condition = self.referrer_condition_combo.get()
        if condition in ("معرف دارد", "بدون معرف"):
            self._set_readonly_text(self.referrer_entry)
        else:
            self.referrer_entry.config(state="normal")
        self.refresh_policy_search_results()
    def change_marketer_combo(self, event=None):
        condition = self.marketer_condition_combo.get()
        if condition in ("بازاریاب دارد", "بدون بازاریاب"):
            self.marketer_combo.set("")
            self.marketer_combo.config(state="disabled")
            self.selected_marketer = None
        else:
            self.marketer_combo.config(state="readonly")
        self.refresh_policy_search_results()
    def select_marketer(self, event=None):
        marketer_name = self.marketer_combo.get()
        self.selected_marketer = next(
            (
                marketer
                for marketer in self.marketers
                if marketer.name == marketer_name
            ),
            None
        )
        self.refresh_policy_search_results()
    # ==================================================
    # حق بیمه
    # ==================================================
    def format_premium(self, event=None):
        if (
            event is not None
            and getattr(event, "keysym", None) in self.IGNORED_KEYS
        ):
            return
        raw = self.premium_entry.get()
        cursor_position = self.premium_entry.index(tk.INSERT)
        digits_before_cursor = sum(
            1 for char in raw[:cursor_position] if char.isdigit()
        )
        # هر کاراکتر غیر عددی حذف می‌شود
        digits = "".join(char for char in raw if char.isdigit())
        formatted = f"{int(digits):,}" if digits else ""
        if formatted != raw:
            self.premium_entry.delete(0, tk.END)
            self.premium_entry.insert(0, formatted)
            # مکان‌نما بعد از همان تعداد رقم قبلی بماند
            if digits_before_cursor == 0:
                new_cursor_position = 0
            else:
                new_cursor_position = len(formatted)
                counted = 0
                for index, char in enumerate(formatted):
                    if char.isdigit():
                        counted += 1
                        if counted == digits_before_cursor:
                            new_cursor_position = index + 1
                            break
            self.premium_entry.icursor(new_cursor_position)
        self.schedule_policy_search()
    # ==================================================
    # حذف همه فیلترها
    # ==================================================
    def clear_search_filters(self):
        self._cancel_pending_search()
        # ---------------- بیمه‌گذار ----------------
        self.selected_customers = []
        self._set_readonly_text(self.customer_search_entry)
        # ---------------- نوع بیمه ----------------
        self.selected_insurance_type = None
        self.main_insurance_type = None
        self._set_readonly_text(self.insurance_type_search_entry)
        # ---------------- شماره بیمه‌نامه ----------------
        self.policy_number_entry.delete(0, tk.END)
        # ---------------- تاریخ‌ها ----------------
        for name in self.DATE_FILTERS:
            combo, entry, entry2 = self._date_widgets(name)
            entry.delete(0, tk.END)
            entry2.delete(0, tk.END)
            combo.current(0)
            entry2.pack_forget()
        # ---------------- معرف (همان حالت اولیه‌ی فرم) ----------------
        self.referrer_condition_combo.current(3)
        self.referrer_entry.config(state="normal")
        self.referrer_entry.delete(0, tk.END)
        # ---------------- بازاریاب (همان حالت اولیه‌ی فرم) ----------------
        self.selected_marketer = None
        self.marketer_condition_combo.current(2)
        self.marketer_combo.config(state="readonly")
        self.marketer_combo.set("")
        # ---------------- حق بیمه ----------------
        self.premium_entry.delete(0, tk.END)
        self.premium_condition_combo.current(0)
        # ---------------- صندوق ----------------
        self.fund_search_entry.delete(0, tk.END)
        self.fund_condition_combo.current(0)
        # ---------------- مرهوناتی و نوع رکورد ----------------
        self.mortgage_search_var.set("")
        self.record_type_search_var.set("")
        # ---------------- حذف تیک همه فیلترها و مخفی کردن فیلدها ----------------
        for variable in self.search_vars.values():
            variable.set(False)
        for frame in self.search_frames.values():
            frame.grid_remove()
        # نمایش همه
        self.filtered_policies = None
        self.load_data()
def open_policy_list(parent, role="کاربر"):
    return PolicyList(parent, role)