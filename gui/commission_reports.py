import tkinter as tk
from tkinter import ttk, messagebox, filedialog

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from services.commission_report_service import (
    get_agent_commission_report,
    get_marketer_commission_report,
    get_policy_commission_report,
)
from services.date_service import (
    convert_shamsi_to_date,
    convert_date_to_shamsi,
)
from services.marketer_service import load_marketers
from services.insurance_type_service import load_root_insurance_types
from gui.insurance_selector import InsuranceSelector
from gui.policy_selector import open_policy_selector
from gui.window_registry import SingleWindowMixin, register_window
from gui.styles import (
    COLORS,
    configure_report_styles,
    finalize_report_window,
)


# =============================================================
# Excel helpers
# =============================================================
def make_border(color="808080"):
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def style_cell(cell, font=None, alignment=None, border=None, fill=None,
               number_format=None):
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border
    if fill is not None:
        cell.fill = fill
    if number_format is not None:
        cell.number_format = number_format


def num(value):
    return value or 0


def date_to_shamsi(value):
    return convert_date_to_shamsi(value) if value else ""


# =============================================================
# Window
# =============================================================
class CommissionReports(SingleWindowMixin, tk.Toplevel):

    def __init__(self, parent, show_menu=True):
        super().__init__(parent)

        configure_report_styles()

        self.configure(bg=COLORS["surface"])
        self.title("گزارش‌های کارمزد")
        self.geometry("600x500")
        self.resizable(False, False)

        if show_menu:
            self.create_widgets()

    # =========================================================
    # Main UI
    # =========================================================
    def create_widgets(self):
        ttk.Label(
            self,
            text="گزارش‌های کارمزد",
            font=("Tahoma", 16, "bold"),
        ).pack(pady=20)

        commission_frame = ttk.LabelFrame(self, text="")
        commission_frame.pack(fill="x", padx=30, pady=10)

        buttons = [
            ("کارمزد نمایندگی در یک دوره", self.open_period_commission_report),
            ("کارمزد بازاریاب در یک دوره", self.open_marketer_commission_report),
            ("کارمزد نمایندگی به تفکیک رشته", self.open_commission_by_type_report),
            ("کارمزد یک بیمه‌نامه", self.open_policy_commission_report),
        ]
        for text, command in buttons:
            ttk.Button(
                commission_frame,
                text=text,
                command=command,
            ).pack(fill="x", padx=15, pady=5, ipady=10)

        finalize_report_window(self)

    # =========================================================
    # Shared UI builders
    # =========================================================
    @staticmethod
    def _filter_row(frame, row, label_text):
        ttk.Label(
            frame,
            text=label_text,
            font=("Tahoma", 10),
        ).grid(row=row, column=1, padx=10, pady=10, sticky="e")

    def _build_filter_header(self, frame):
        frame.columnconfigure(0, weight=1)
        ttk.Label(
            frame,
            text="فیلتر گزارش",
            font=("Tahoma", 10, "bold"),
        ).grid(
            row=0, column=0, columnspan=2,
            padx=15, pady=(5, 10), sticky="e",
        )

    def _build_basis_combo(self, frame, row):
        self._filter_row(frame, row, "مبنا:")
        var = tk.StringVar(value="حق بیمه")
        combo = ttk.Combobox(
            frame,
            textvariable=var,
            values=["حق بیمه", "وصولی"],
            state="readonly",
            width=18,
            justify="right",
        )
        combo.grid(row=row, column=0, padx=10, pady=10, sticky="e")
        return var, combo

    def _build_date_entry(self, frame, row, label_text):
        self._filter_row(frame, row, label_text)
        var = tk.StringVar()
        entry = ttk.Entry(
            frame,
            textvariable=var,
            justify="center",
            width=20,
        )
        entry.grid(row=row, column=0, padx=10, pady=10, sticky="e")
        return var, entry

    @staticmethod
    def _build_summary(window, items):
        """
        items: list of (title, key)
        returns: (frame, vars_by_key, labels_by_key)
        """
        frame = ttk.LabelFrame(window)
        frame.pack(fill="x", padx=30, pady=10)

        ttk.Label(
            frame,
            text="خلاصه گزارش",
            font=("Tahoma", 10, "bold"),
            anchor="e",
        ).pack(fill="x", padx=15, pady=(8, 12))

        value_vars = {}
        title_labels = {}

        for title, key in items:
            row_frame = ttk.Frame(frame)
            row_frame.pack(fill="x", padx=15, pady=4)

            value_var = tk.StringVar(value="0")
            value_vars[key] = value_var

            title_label = ttk.Label(
                row_frame,
                text=f"{title}:",
                font=("Tahoma", 10),
                anchor="e",
            )
            title_label.pack(side="right")
            title_labels[key] = title_label

            ttk.Label(
                row_frame,
                textvariable=value_var,
                font=("Tahoma", 10, "bold"),
                anchor="e",
            ).pack(side="right", padx=(0, 10))

        return frame, value_vars, title_labels

    @staticmethod
    def _parse_dates(start_shamsi, end_shamsi, parent):
        """برمی‌گرداند (start_date, end_date) یا None در صورت خطا."""
        if not start_shamsi or not end_shamsi:
            messagebox.showwarning(
                "خطا",
                "لطفاً تاریخ شروع و پایان را وارد کنید.",
                parent=parent,
            )
            return None

        try:
            start_date = convert_shamsi_to_date(start_shamsi)
            end_date = convert_shamsi_to_date(end_shamsi)
        except Exception:
            messagebox.showerror(
                "خطا",
                "فرمت تاریخ صحیح نیست.",
                parent=parent,
            )
            return None

        if start_date > end_date:
            messagebox.showwarning(
                "خطا",
                "تاریخ شروع نمی‌تواند بعد از تاریخ پایان باشد.",
                parent=parent,
            )
            return None

        return start_date, end_date

    @staticmethod
    def _period_window_key(insurance_type_id, issue_code):
        """گزارش دوره‌ای، به تفکیک رشته و بیمه‌نامه‌ای، هر کدام پنجره‌ی جدا دارند."""
        if issue_code is not None:
            return f"commission_policy_period_{issue_code}"
        if insurance_type_id is not None:
            return "commission_by_type"
        return "commission_period"

    # =========================================================
    # Period commission report window
    # =========================================================
    def open_period_commission_report(
        self,
        insurance_type_id=None,
        insurance_type_name=None,
        issue_code=None,
        policy_number=None,
    ):
        window_key = self._period_window_key(insurance_type_id, issue_code)

        if self.reuse_window(window_key):
            return

        window = tk.Toplevel(self)
        register_window(window_key, window)
        window.configure(bg=COLORS["surface"])

        if insurance_type_id is None:
            window.title("گزارش کارمزد نمایندگی")
        else:
            window.title(f"گزارش کارمزد نمایندگی - {insurance_type_name}")

        window.geometry("650x780")
        window.resizable(False, False)

        # اطلاعات گزارش
        self.current_period_insurance_type_id = insurance_type_id
        self.current_period_insurance_type_name = insurance_type_name or ""
        self.current_period_issue_code = issue_code
        self.current_period_policy_number = policy_number

        # Title
        if issue_code is not None:
            report_title = f"گزارش کارمزد بیمه‌نامه {policy_number}"
        elif insurance_type_id is not None:
            report_title = f"گزارش کارمزد نمایندگی - {insurance_type_name}"
        else:
            report_title = "گزارش کارمزد نمایندگی"

        ttk.Label(
            window,
            text=report_title,
            style="CommissionTitle.TLabel",
        ).pack(pady=(20, 5), anchor="n", padx=30)

        # Filter frame
        filter_frame = ttk.LabelFrame(window, style="Commission.TLabelframe")
        filter_frame.pack(fill="x", padx=30, pady=15)
        self._build_filter_header(filter_frame)

        self.period_basis_var, self.period_basis_combo = (
            self._build_basis_combo(filter_frame, 1)
        )
        self.period_start_var, self.period_start_entry = (
            self._build_date_entry(filter_frame, 2, "از تاریخ:")
        )
        self.period_end_var, self.period_end_entry = (
            self._build_date_entry(filter_frame, 3, "تا تاریخ:")
        )

        ttk.Button(
            filter_frame,
            text="نمایش گزارش",
            command=self.show_period_commission_report,
        ).grid(row=4, column=0, columnspan=2, padx=10, pady=15)

        # Summary
        summary_items = [
            ("مبلغ ناخالص حق بیمه", "gross_amount"),
            ("مبلغ خالص حق بیمه", "net_amount"),
            ("کارمزد ناخالص", "gross_commission"),
            ("کارمزد خالص", "net_commission"),
            ("هزینه صدور ناخالص", "gross_issuance"),
            ("هزینه صدور خالص", "net_issuance"),
            ("مجموع کارمزد و هزینه صدور ناخالص", "total_commission_issuance"),
            ("سهم بازاریاب", "marketer_commission"),
            ("تعداد رکورد", "record_count"),
        ]
        (
            self.period_summary_frame,
            self.period_summary_vars,
            self.period_summary_labels,
        ) = self._build_summary(window, summary_items)

        # Excel
        ttk.Button(
            window,
            text="مشاهده جزئیات کارمزد از طریق فایل اکسل",
            command=self.export_period_commission_report,
        ).pack(pady=5, ipadx=15, ipady=10)

        finalize_report_window(window)
        self.period_report_window = window

    # =========================================================
    # Marketer commission report window
    # =========================================================
    def open_marketer_commission_report(self):
        if self.reuse_window("commission_marketer"):
            return

        window = tk.Toplevel(self)
        register_window("commission_marketer", window)
        window.configure(bg=COLORS["surface"])
        window.title("گزارش کارمزد بازاریاب")
        window.geometry("650x720")
        window.resizable(False, False)

        ttk.Label(
            window,
            text="گزارش کارمزد بازاریاب",
            font=("Tahoma", 16, "bold"),
        ).pack(pady=(20, 5), anchor="n", padx=30)

        # Filter frame
        filter_frame = ttk.LabelFrame(window)
        filter_frame.pack(fill="x", padx=30, pady=15)
        self._build_filter_header(filter_frame)

        # Marketer
        self._filter_row(filter_frame, 1, "بازاریاب:")

        self.marketer_period_marketers = load_marketers()
        marketer_names = [m.name for m in self.marketer_period_marketers]

        self.marketer_period_marketer_var = tk.StringVar()
        self.marketer_period_marketer_combo = ttk.Combobox(
            filter_frame,
            textvariable=self.marketer_period_marketer_var,
            values=marketer_names,
            state="readonly",
            width=18,
            justify="right",
        )
        self.marketer_period_marketer_combo.grid(
            row=1, column=0, padx=10, pady=10, sticky="e"
        )
        if marketer_names:
            self.marketer_period_marketer_combo.current(0)

        self.marketer_period_basis_var, self.marketer_period_basis_combo = (
            self._build_basis_combo(filter_frame, 2)
        )
        self.marketer_period_start_var, self.marketer_period_start_entry = (
            self._build_date_entry(filter_frame, 3, "از تاریخ:")
        )
        self.marketer_period_end_var, self.marketer_period_end_entry = (
            self._build_date_entry(filter_frame, 4, "تا تاریخ:")
        )

        ttk.Button(
            filter_frame,
            text="نمایش گزارش",
            command=self.show_marketer_commission_report,
        ).grid(row=5, column=0, columnspan=2, padx=10, pady=15)

        # Summary
        summary_items = [
            ("مبلغ ناخالص", "gross_amount"),
            ("مبلغ خالص", "net_amount"),
            ("کارمزد ناخالص", "gross_commission"),
            ("کارمزد خالص", "net_commission"),
            ("سهم بازاریاب", "marketer_commission"),
            ("تعداد رکورد", "record_count"),
        ]
        (
            self.marketer_period_summary_frame,
            self.marketer_period_summary_vars,
            _,
        ) = self._build_summary(window, summary_items)

        # Excel
        ttk.Button(
            window,
            text="مشاهده جزئیات کارمزد از طریق فایل اکسل",
            command=self.export_marketer_commission_report,
        ).pack(pady=5, ipadx=15, ipady=10)

        finalize_report_window(window)
        self.marketer_period_report_window = window

    # =========================================================
    # Show period report
    # =========================================================
    def show_period_commission_report(self):
        parent = self.period_report_window

        dates = self._parse_dates(
            self.period_start_var.get().strip(),
            self.period_end_var.get().strip(),
            parent,
        )
        if dates is None:
            return
        start_date, end_date = dates

        basis = (
            "premium"
            if self.period_basis_var.get() == "حق بیمه"
            else "collection"
        )

        try:
            report = get_agent_commission_report(
                basis=basis,
                start_date=start_date,
                end_date=end_date,
                insurance_type_id=self.current_period_insurance_type_id,
                issue_code=self.current_period_issue_code,
            )
        except Exception as e:
            messagebox.showerror(
                "خطا",
                f"خطا در تهیه گزارش:\n{e}",
                parent=parent,
            )
            return

        details = report.get("details", [])

        # مبلغ ناخالص / خالص و عنوان‌ها بر اساس مبنا
        if basis == "premium":
            gross_amount = report.get("total_premium", 0)
            net_amount = report.get("total_net_premium", 0)
            gross_title = "مبلغ ناخالص حق بیمه"
            net_title = "مبلغ خالص حق بیمه"
        else:
            gross_amount = report.get("total_collection", 0)
            net_amount = report.get("total_net_collection", 0)
            gross_title = "مبلغ ناخالص وصولی"
            net_title = "مبلغ خالص وصولی"

        self.period_summary_labels["gross_amount"].configure(
            text=f"{gross_title}:"
        )
        self.period_summary_labels["net_amount"].configure(
            text=f"{net_title}:"
        )

        gross_commission = report.get("total_gross_commission", 0)
        gross_issuance = report.get("total_gross_issuance", 0)

        values = {
            "gross_amount": gross_amount,
            "net_amount": net_amount,
            "gross_commission": gross_commission,
            "net_commission": report.get("total_commission", 0),
            "gross_issuance": gross_issuance,
            "net_issuance": report.get("total_net_issuance", 0),
            "total_commission_issuance": gross_commission + gross_issuance,
            "marketer_commission": report.get("total_marketer_commission", 0),
            "record_count": len(details),
        }
        for key, value in values.items():
            self.period_summary_vars[key].set(f"{value:,}")

        # نگهداری گزارش برای Excel
        self.current_period_report = report
        self.current_period_report_basis = basis
        self.current_period_start_date = start_date
        self.current_period_end_date = end_date

        messagebox.showinfo(
            "گزارش",
            "گزارش با موفقیت تهیه شد.",
            parent=parent,
        )

    # =========================================================
    # Show marketer report
    # =========================================================
    def show_marketer_commission_report(self):
        parent = self.marketer_period_report_window

        if not self.marketer_period_marketers:
            messagebox.showwarning(
                "گزارش",
                "هیچ بازاریاب فعالی وجود ندارد.",
                parent=parent,
            )
            return

        selected_index = self.marketer_period_marketer_combo.current()
        if selected_index < 0:
            messagebox.showwarning(
                "خطا",
                "لطفاً بازاریاب را انتخاب کنید.",
                parent=parent,
            )
            return

        marketer = self.marketer_period_marketers[selected_index]

        dates = self._parse_dates(
            self.marketer_period_start_var.get().strip(),
            self.marketer_period_end_var.get().strip(),
            parent,
        )
        if dates is None:
            return
        start_date, end_date = dates

        basis = (
            "premium"
            if self.marketer_period_basis_var.get() == "حق بیمه"
            else "collection"
        )

        try:
            report = get_marketer_commission_report(
                marketer_id=marketer.id,
                basis=basis,
                start_date=start_date,
                end_date=end_date,
            )
        except Exception as e:
            messagebox.showerror(
                "خطا",
                f"خطا در تهیه گزارش:\n{e}",
                parent=parent,
            )
            return

        # مرتب‌سازی تاریخ از کم به زیاد
        details = sorted(
            report.get("details", []),
            key=lambda item: (
                item.get("date") or start_date,
                item.get("policy_id") or 0,
            ),
        )
        report["details"] = details

        # مبلغ ناخالص / خالص
        if basis == "premium":
            gross_amount = report.get("total_premium", 0)
            net_amount = report.get("total_net_premium", 0)
        else:
            gross_amount = sum(
                num(d.get("collection_amount", 0)) for d in details
            )
            net_amount = sum(
                num(d.get("net_collection_amount", 0)) for d in details
            )

        values = {
            "gross_amount": gross_amount,
            "net_amount": net_amount,
            "gross_commission": report.get("total_gross_commission", 0),
            "net_commission": report.get("total_commission", 0),
            "marketer_commission": report.get("total_marketer_commission", 0),
            "record_count": len(details),
        }
        for key, value in values.items():
            self.marketer_period_summary_vars[key].set(f"{value:,}")

        # نگهداری گزارش برای Excel
        self.current_marketer_period_report = report
        self.current_marketer_period_marketer_id = marketer.id
        self.current_marketer_period_marketer_name = marketer.name
        self.current_marketer_period_basis = basis
        self.current_marketer_period_start_date = start_date
        self.current_marketer_period_end_date = end_date

        messagebox.showinfo(
            "گزارش",
            "گزارش با موفقیت تهیه شد.",
            parent=parent,
        )

    # =========================================================
    # Excel export - period (agency)
    # =========================================================
    def export_period_commission_report(self):
        parent = self.period_report_window

        if not hasattr(self, "current_period_report"):
            messagebox.showwarning(
                "گزارش",
                "ابتدا گزارش را نمایش دهید.",
                parent=parent,
            )
            return

        file_path = filedialog.asksaveasfilename(
            parent=parent,
            title="ذخیره گزارش کارمزد",
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            initialfile="گزارش_کارمزد_نمایندگی.xlsx",
        )
        if not file_path:
            return

        try:
            report = self.current_period_report
            basis = self.current_period_report_basis
            start_date = self.current_period_start_date
            end_date = self.current_period_end_date
            details = report.get("details", [])

            # عنوان مبلغ بر اساس مبنا
            if basis == "premium":
                title_gross = "حق بیمه ناخالص"
                title_net = "حق بیمه خالص"
                total_basis_gross = report.get("total_premium", 0)
                total_basis_net = report.get("total_net_premium", 0)
            else:
                title_gross = "وصولی ناخالص"
                title_net = "وصولی خالص"
                total_basis_gross = report.get("total_collection", 0)
                total_basis_net = report.get("total_net_collection", 0)

            # همه مقادیر از report گرفته می‌شوند تا دقیقاً همان چیزی
            # باشند که UI نمایش می‌دهد.
            total_net_issuance = report.get("total_net_issuance", 0)
            total_gross_issuance = report.get("total_gross_issuance", 0)
            total_commission = report.get("total_commission", 0)
            total_gross_commission = report.get("total_gross_commission", 0)
            total_marketer_commission = report.get(
                "total_marketer_commission", 0
            )
            total_gross_commission_and_issuance = (
                total_gross_commission + total_gross_issuance
            )

            # ---------------------------------------------------------
            # Styles
            # ---------------------------------------------------------
            title_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
            header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
            summary_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")

            white_font = Font(name="Tahoma", size=12, bold=True, color="FFFFFF")
            header_font = Font(name="Tahoma", size=10, bold=True)
            normal_font = Font(name="Tahoma", size=10)

            border = make_border("B7B7B7")
            centered = Alignment(horizontal="center", vertical="center")
            right_aligned = Alignment(horizontal="right", vertical="center")

            workbook = Workbook()

            # =========================================================
            # Sheet 1 - خلاصه
            # =========================================================
            summary_sheet = workbook.active
            summary_sheet.title = "خلاصه"
            summary_sheet.sheet_view.rightToLeft = True

            summary_sheet.merge_cells("A1:C1")
            title_cell = summary_sheet["A1"]
            title_cell.value = "گزارش کارمزد نمایندگی"
            style_cell(title_cell, white_font, centered, fill=title_fill)
            summary_sheet.row_dimensions[1].height = 28

            # مشخصات گزارش
            summary_sheet["A3"] = "عنوان"
            summary_sheet["B3"] = "مقدار"
            for cell in (summary_sheet["A3"], summary_sheet["B3"]):
                style_cell(cell, header_font, centered, border, header_fill)

            info_rows = [
                ("مبنا", "حق بیمه" if basis == "premium" else "وصولی"),
                ("از تاریخ", convert_date_to_shamsi(start_date)),
                ("تا تاریخ", convert_date_to_shamsi(end_date)),
                ("تعداد رکورد", len(details)),
            ]
            row = 4
            for title, value in info_rows:
                summary_sheet.cell(row=row, column=1, value=title)
                summary_sheet.cell(row=row, column=2, value=value)
                for col in (1, 2):
                    style_cell(
                        summary_sheet.cell(row=row, column=col),
                        normal_font, right_aligned, border,
                    )
                row += 1

            # جدول اصلی خلاصه
            row += 1
            for col, text in enumerate(
                ["شاخص", "مقدار گزارش", "جمع از جزئیات Excel"], start=1
            ):
                cell = summary_sheet.cell(row=row, column=col, value=text)
                style_cell(cell, header_font, centered, border, summary_fill)
            summary_header_row = row

            # =========================================================
            # Sheet 2 - جزئیات
            # =========================================================
            detail_headers = [
                "ردیف",
                "بیمه‌گذار",
                "رشته بیمه",
                "ش بیمه",
                "ش الحاقیه",
                "تاریخ",
                title_gross,
                title_net,
                "کارمزد خالص",
                "هزینه صدور خالص",
                "جمع کارمزد و هزینه صدور ناخالص",
                "سهم بازاریاب",
            ]

            detail_sheet = workbook.create_sheet("جزئیات")
            detail_sheet.sheet_view.rightToLeft = True
            detail_sheet.freeze_panes = "A2"

            for col, header in enumerate(detail_headers, start=1):
                cell = detail_sheet.cell(row=1, column=col, value=header)
                style_cell(cell, header_font, centered, border, header_fill)

            for index, detail in enumerate(details, start=1):
                if basis == "premium":
                    amount_gross = detail.get("premium", 0)
                    amount_net = detail.get("net_premium", 0)
                else:
                    amount_gross = detail.get("collection_amount", 0)
                    amount_net = detail.get("net_collection_amount", 0)

                # جمع فقط برای همین ردیف
                row_total_gross = (
                    num(detail.get("gross_commission", 0))
                    + num(detail.get("gross_issuance", 0))
                )

                row_values = [
                    index,
                    detail.get("customer_name", ""),
                    detail.get("insurance_type_name", ""),
                    detail.get("policy_number", ""),
                    detail.get("endorsement_number", ""),
                    date_to_shamsi(detail.get("date")),
                    amount_gross,
                    amount_net,
                    detail.get("commission", 0),
                    detail.get("net_issuance", 0),
                    row_total_gross,
                    detail.get("marketer_commission", 0),
                ]

                for col, value in enumerate(row_values, start=1):
                    cell = detail_sheet.cell(
                        row=index + 1, column=col, value=value
                    )
                    style_cell(cell, normal_font, right_aligned, border)
                    if col in range(7, 13):  # ستون‌های مبلغ
                        cell.number_format = "#,##0"

            detail_last_row = max(len(details) + 1, 2)

            # =========================================================
            # فرمول‌های خلاصه
            # (عنوان، مقدار گزارش، ستون جزئیات یا None)
            #
            # برای «کارمزد ناخالص» و «هزینه صدور ناخالص» ستون مجزا
            # در شیت جزئیات وجود ندارد؛ مقدار گزارش عیناً درج می‌شود.
            # =========================================================
            metrics = [
                (title_gross, total_basis_gross, "G"),
                (title_net, total_basis_net, "H"),
                ("کارمزد خالص", total_commission, "I"),
                ("کارمزد ناخالص", total_gross_commission, None),
                ("هزینه صدور خالص", total_net_issuance, "J"),
                ("هزینه صدور ناخالص", total_gross_issuance, None),
                (
                    "جمع کارمزد و هزینه صدور ناخالص",
                    total_gross_commission_and_issuance,
                    "K",
                ),
                ("سهم بازاریاب", total_marketer_commission, "L"),
            ]

            summary_row = summary_header_row + 1
            for metric_name, report_value, column_letter in metrics:
                summary_sheet.cell(row=summary_row, column=1, value=metric_name)
                summary_sheet.cell(row=summary_row, column=2, value=report_value)

                if column_letter:
                    third_value = (
                        f"=SUM('جزئیات'!${column_letter}$2:"
                        f"${column_letter}${detail_last_row})"
                    )
                else:
                    third_value = report_value
                summary_sheet.cell(row=summary_row, column=3, value=third_value)

                for col in (1, 2, 3):
                    cell = summary_sheet.cell(row=summary_row, column=col)
                    style_cell(cell, normal_font, right_aligned, border)
                    if col in (2, 3):
                        cell.number_format = "#,##0"
                summary_row += 1

            # تعداد رکورد
            summary_sheet.cell(row=summary_row, column=1, value="تعداد رکورد")
            summary_sheet.cell(row=summary_row, column=2, value=len(details))
            summary_sheet.cell(
                row=summary_row,
                column=3,
                value=f"=COUNTA('جزئیات'!$A$2:$A${detail_last_row})",
            )
            for col in (1, 2, 3):
                cell = summary_sheet.cell(row=summary_row, column=col)
                style_cell(cell, normal_font, right_aligned, border)
                if col in (2, 3):
                    cell.number_format = "#,##0"

            # عرض ستون‌ها
            for letter, width in {"A": 42, "B": 24, "C": 27}.items():
                summary_sheet.column_dimensions[letter].width = width

            detail_widths = {
                1: 5, 2: 30, 3: 22, 4: 8, 5: 8, 6: 12,
                7: 15, 8: 15, 9: 15, 10: 15, 11: 15, 12: 15,
            }
            for col, width in detail_widths.items():
                detail_sheet.column_dimensions[
                    get_column_letter(col)
                ].width = width
            detail_sheet.row_dimensions[1].height = 32

            # محاسبه مجدد فرمول‌ها هنگام باز شدن فایل
            try:
                workbook.calculation.fullCalcOnLoad = True
                workbook.calculation.forceFullCalc = True
                workbook.calculation.calcMode = "auto"
            except Exception:
                pass

            workbook.save(file_path)

            messagebox.showinfo(
                "گزارش",
                "فایل Excel با موفقیت ایجاد شد.",
                parent=parent,
            )

        except Exception as e:
            messagebox.showerror(
                "خطا",
                f"خطا در ایجاد فایل اکسل:\n{e}",
                parent=parent,
            )

    # =========================================================
    # Excel export - marketer
    # =========================================================
    def export_marketer_commission_report(self):
        parent = self.marketer_period_report_window

        if not hasattr(self, "current_marketer_period_report"):
            messagebox.showwarning(
                "گزارش",
                "ابتدا گزارش را نمایش دهید.",
                parent=parent,
            )
            return

        report = self.current_marketer_period_report
        basis = self.current_marketer_period_basis
        start_date = self.current_marketer_period_start_date
        end_date = self.current_marketer_period_end_date
        marketer_name = self.current_marketer_period_marketer_name

        # مرتب‌سازی تاریخ از کم به زیاد
        details = sorted(
            report.get("details", []),
            key=lambda item: (
                item.get("date") or start_date,
                item.get("policy_id") or 0,
            ),
        )

        # عنوان ستون‌ها و جمع‌ها
        if basis == "collection":
            premium_header = "حق بیمه وصولی"
            net_premium_header = "حق بیمه وصولی خالص"
            date_header = "تاریخ وصولی"
            total_premium = sum(
                num(d.get("collection_amount", 0)) for d in details
            )
            total_net_premium = sum(
                num(d.get("net_collection_amount", 0)) for d in details
            )
        else:
            premium_header = "حق بیمه کل"
            net_premium_header = "حق بیمه خالص"
            date_header = "تاریخ صدور"
            total_premium = report.get("total_premium", 0)
            total_net_premium = report.get("total_net_premium", 0)

        headers = [
            "ردیف",
            "بیمه گذار",
            "رشته بیمه",
            premium_header,
            net_premium_header,
            "کل کارمزد",
            "کارمزد سهم معرف",
            "کارمزد رشته",
            date_header,
        ]

        # نام پیش‌فرض فایل
        end_date_shamsi = convert_date_to_shamsi(end_date)
        default_file_name = (
            f"{marketer_name}{end_date_shamsi.replace('/', '')}.xlsx"
        )

        file_path = filedialog.asksaveasfilename(
            parent=parent,
            title="ذخیره گزارش کارمزد بازاریاب",
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            initialfile=default_file_name,
        )
        if not file_path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "گزارش کارمزد"
            ws.sheet_view.rightToLeft = True

            # Header
            for col, header in enumerate(headers, start=1):
                ws.cell(row=1, column=col, value=header)

            # Detail rows
            row_number = 2
            for detail in details:
                if basis == "collection":
                    premium = num(detail.get("collection_amount", 0))
                    net_premium = num(detail.get("net_collection_amount", 0))
                else:
                    premium = num(detail.get("premium", 0))
                    net_premium = num(detail.get("net_premium", 0))

                row_values = [
                    row_number - 1,
                    detail.get("customer_name", ""),
                    detail.get("insurance_type_name", ""),
                    premium,
                    net_premium,
                    num(detail.get("commission", 0)),
                    num(detail.get("marketer_commission", 0)),
                    num(detail.get("commission_percent", 0)),
                    date_to_shamsi(detail.get("date")),
                ]
                for col, value in enumerate(row_values, start=1):
                    ws.cell(row=row_number, column=col, value=value)
                row_number += 1

            # Total row
            total_row = row_number
            ws.merge_cells(
                start_row=total_row, start_column=1,
                end_row=total_row, end_column=3,
            )
            ws.cell(row=total_row, column=1, value="جمع کل")
            ws.cell(row=total_row, column=4, value=total_premium)
            ws.cell(row=total_row, column=5, value=total_net_premium)
            ws.cell(
                row=total_row, column=6,
                value=report.get("total_commission", 0),
            )
            ws.cell(
                row=total_row, column=7,
                value=report.get("total_marketer_commission", 0),
            )

            # رنگ و فونت جمع کل
            total_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
            total_font = Font(bold=True)
            centered = Alignment(horizontal="center", vertical="center")

            for col in range(1, 8):
                style_cell(
                    ws.cell(row=total_row, column=col),
                    total_font, centered, fill=total_fill,
                )

            # رنگ و فونت Header
            header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
            header_font = Font(bold=True, color="FFFFFF")
            for col in range(1, 10):
                style_cell(
                    ws.cell(row=1, column=col),
                    header_font, centered, fill=header_fill,
                )

            # Border
            border = make_border("808080")
            for r in ws.iter_rows(
                min_row=1, max_row=total_row, min_col=1, max_col=9
            ):
                for cell in r:
                    cell.border = border

            # ستون 8 و 9 در ردیف جمع کل: بدون رنگ و Border
            for col in (8, 9):
                cell = ws.cell(row=total_row, column=col)
                cell.border = Border()
                cell.fill = PatternFill(fill_type=None)

            # Number format
            for r in range(2, total_row + 1):
                for col in (4, 5, 6, 7):
                    ws.cell(row=r, column=col).number_format = "#,##0"
            for r in range(2, total_row):
                ws.cell(row=r, column=8).number_format = "0.0#"

            # Column widths
            column_widths = {
                1: 6, 2: 35, 3: 25, 4: 15, 5: 18,
                6: 15, 7: 15, 8: 10, 9: 15,
            }
            for col, width in column_widths.items():
                ws.column_dimensions[get_column_letter(col)].width = width

            ws.freeze_panes = "A2"

            wb.save(file_path)

            messagebox.showinfo(
                "گزارش",
                "فایل Excel با موفقیت ایجاد شد.",
                parent=parent,
            )

        except Exception as e:
            messagebox.showerror(
                "خطا",
                f"خطا در ایجاد فایل اکسل:\n{e}",
                parent=parent,
            )

    # =========================================================
    # Other reports
    # =========================================================
    def open_commission_by_type_report(self):
        if self.reuse_window("commission_by_type"):
            return

        root_types = load_root_insurance_types()
        InsuranceSelector(
            self,
            root_types,
            self.insurance_type_selected_for_report,
        )

    def insurance_type_selected_for_report(self, selected_path):
        insurance_type = selected_path[-1]
        self.open_period_commission_report(
            insurance_type_id=insurance_type.id,
            insurance_type_name=" > ".join(item.name for item in selected_path),
        )

    def open_policy_commission_report(self):
        if self.reuse_window("commission_policy"):
            return

        open_policy_selector(
            self,
            self.policy_selected_for_commission_report,
        )

    def policy_selected_for_commission_report(self, policy):
        self.open_single_policy_commission_report(policy)

    # =========================================================
    # Single policy report
    # =========================================================
    def open_single_policy_commission_report(self, policy):
        window = tk.Toplevel(self)
        window.configure(bg=COLORS["surface"])
        window.title(
            f"گزارش کارمزد بیمه‌نامه - کد رایانه بیمه: {policy.issue_code}"
        )
        window.geometry("700x700")
        window.resizable(False, False)

        try:
            report = get_policy_commission_report(
                issue_code=policy.issue_code
            )
        except Exception as e:
            messagebox.showerror(
                "خطا",
                f"خطا در تهیه گزارش:\n{e}",
                parent=window,
            )
            window.destroy()
            return

        if not report:
            messagebox.showwarning(
                "گزارش",
                "برای این بیمه‌نامه گزارشی پیدا نشد.",
                parent=window,
            )
            window.destroy()
            return

        self.current_policy_commission_report = report
        self.single_policy_report_window = window
        register_window("commission_policy", window)

        # عنوان
        ttk.Label(
            window,
            text=(
                "گزارش کارمزد بیمه‌نامه - "
                f"کد رایانه بیمه: {report['issue_code']}"
            ),
            font=("Tahoma", 16, "bold"),
        ).pack(pady=(20, 15))

        # Summary
        summary_frame = ttk.LabelFrame(window, text="خلاصه گزارش")
        summary_frame.pack(fill="x", padx=30, pady=10)

        summary_items = [
            ("مبلغ خالص حق بیمه:", report["original_net_premium"]),
            ("مبلغ ناخالص حق بیمه:", report["original_gross_premium"]),
            ("مبلغ خالص الحاقیه ها:", report["endorsement_net_premium"]),
            ("مبلغ ناخالص الحاقیه ها:", report["endorsement_gross_premium"]),
            ("مجموع مبلغ وصول شده:", report["total_collection"]),
            ("مبلغ باقی مانده وصول نشده:", report["remaining_uncollected"]),
            (
                "کارمزد خالص کل بیمه نامه و الحاقیه ها بر مبنای حق بیمه:",
                report["premium_net_commission"],
            ),
            (
                "کارمزد ناخالص کل بیمه نامه و الحاقیه ها بر مبنای حق بیمه:",
                report["premium_gross_commission"],
            ),
            (
                "هزینه صدور خالص کل بیمه نامه و الحاقیه ها بر مبنای حق بیمه:",
                report["premium_net_issuance"],
            ),
            (
                "هزینه صدور ناخالص کل بیمه نامه و الحاقیه ها بر مبنای حق بیمه:",
                report["premium_gross_issuance"],
            ),
            (
                "کارمزد خالص کل بیمه نامه و الحاقیه ها بر مبنای وصولی:",
                report["collection_net_commission"],
            ),
            (
                "کارمزد ناخالص کل بیمه نامه و الحاقیه ها بر مبنای وصولی:",
                report["collection_gross_commission"],
            ),
            (
                "هزینه صدور خالص کل بیمه نامه و الحاقیه ها بر مبنای وصولی:",
                report["collection_net_issuance"],
            ),
            (
                "هزینه صدور ناخالص کل بیمه نامه و الحاقیه ها بر مبنای وصولی:",
                report["collection_gross_issuance"],
            ),
            ("سهم بازاریاب بر مبنای حق بیمه:", report["premium_marketer_share"]),
            ("سهم بازاریاب بر مبنای وصولی:", report["collection_marketer_share"]),
        ]

        for title, value in summary_items:
            row_frame = ttk.Frame(summary_frame)
            row_frame.pack(fill="x", padx=15, pady=4)

            ttk.Label(
                row_frame,
                text=title,
                font=("Tahoma", 10),
                anchor="e",
                justify="right",
                wraplength=500,
            ).pack(side="right")

            ttk.Label(
                row_frame,
                text=f"{value:,}",
                font=("Tahoma", 10, "bold"),
                anchor="e",
            ).pack(side="left", padx=10)

        # Excel
        ttk.Button(
            window,
            text="مشاهده جزئیات از طریق فایل اکسل",
            command=self.export_single_policy_commission_report,
        ).pack(pady=20, ipadx=20, ipady=10)

        finalize_report_window(window)

    def export_single_policy_commission_report(self):
        parent = self.single_policy_report_window

        if not hasattr(self, "current_policy_commission_report"):
            messagebox.showwarning(
                "گزارش",
                "ابتدا گزارش را تهیه کنید.",
                parent=parent,
            )
            return

        report = self.current_policy_commission_report

        file_path = filedialog.asksaveasfilename(
            parent=parent,
            title="ذخیره گزارش کارمزد بیمه‌نامه",
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            initialfile=f"{report['issue_code']}.xlsx",
        )
        if not file_path:
            return

        try:
            workbook = Workbook()

            # ---------------------------------------------------------
            # Styles
            # ---------------------------------------------------------
            header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
            total_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")

            header_font = Font(name="Tahoma", size=10, bold=True)
            total_font = Font(name="Tahoma", size=10, bold=True)
            normal_font = Font(name="Tahoma", size=10)

            border = make_border("808080")
            centered = Alignment(horizontal="center", vertical="center")
            right_aligned = Alignment(horizontal="right", vertical="center")

            column_widths = {
                1: 6, 2: 30, 3: 22, 4: 15, 5: 12, 6: 12,
                7: 16, 8: 16, 9: 16, 10: 16, 11: 24, 12: 16,
            }
            amount_columns = range(7, 13)

            # ---------------------------------------------------------
            # ساخت یک Sheet جزئیات
            # ---------------------------------------------------------
            def fill_detail_sheet(sheet, details, gross_title, net_title,
                                  gross_key, net_key):
                sheet.sheet_view.rightToLeft = True
                sheet.freeze_panes = "A2"

                headers = [
                    "ردیف",
                    "بیمه‌گذار",
                    "رشته بیمه",
                    "شماره بیمه",
                    "شماره الحاقیه",
                    "تاریخ",
                    gross_title,
                    net_title,
                    "کارمزد خالص",
                    "هزینه صدور خالص",
                    "جمع کارمزد و هزینه صدور ناخالص",
                    "سهم بازاریاب",
                ]

                for col, header in enumerate(headers, start=1):
                    cell = sheet.cell(row=1, column=col, value=header)
                    style_cell(cell, header_font, centered, border, header_fill)
                sheet.row_dimensions[1].height = 32

                # Detail rows
                for index, detail in enumerate(details, start=1):
                    total_gross = (
                        num(detail.get("gross_commission", 0))
                        + num(detail.get("gross_issuance", 0))
                    )
                    values = [
                        index,
                        detail.get("customer_name", ""),
                        detail.get("insurance_type_name", ""),
                        detail.get("policy_number", ""),
                        detail.get("endorsement_number", 0),
                        date_to_shamsi(detail.get("date")),
                        num(detail.get(gross_key, 0)),
                        num(detail.get(net_key, 0)),
                        num(detail.get("commission", 0)),
                        num(detail.get("net_issuance", 0)),
                        total_gross,
                        num(detail.get("marketer_commission", 0)),
                    ]
                    for col, value in enumerate(values, start=1):
                        cell = sheet.cell(
                            row=index + 1, column=col, value=value
                        )
                        style_cell(cell, normal_font, right_aligned, border)
                        if col in amount_columns:
                            cell.number_format = "#,##0"

                # Total row
                total_row = len(details) + 2
                sheet.merge_cells(
                    start_row=total_row, start_column=1,
                    end_row=total_row, end_column=6,
                )
                sheet.cell(row=total_row, column=1, value="جمع کل")

                total_values = {
                    7: sum(num(d.get(gross_key, 0)) for d in details),
                    8: sum(num(d.get(net_key, 0)) for d in details),
                    9: sum(num(d.get("commission", 0)) for d in details),
                    10: sum(num(d.get("net_issuance", 0)) for d in details),
                    11: sum(
                        num(d.get("gross_commission", 0))
                        + num(d.get("gross_issuance", 0))
                        for d in details
                    ),
                    12: sum(
                        num(d.get("marketer_commission", 0)) for d in details
                    ),
                }
                for col, value in total_values.items():
                    sheet.cell(row=total_row, column=col, value=value)

                for col in range(1, 13):
                    cell = sheet.cell(row=total_row, column=col)
                    style_cell(cell, total_font, centered, border, total_fill)
                    if col in amount_columns:
                        cell.number_format = "#,##0"

                for col, width in column_widths.items():
                    sheet.column_dimensions[
                        get_column_letter(col)
                    ].width = width

            # Sheet اول: حق بیمه (از Sheet پیش‌فرض استفاده می‌کنیم)
            premium_sheet = workbook.active
            premium_sheet.title = "حق بیمه"
            fill_detail_sheet(
                sheet=premium_sheet,
                details=report.get("premium_details", []),
                gross_title="حق بیمه ناخالص",
                net_title="حق بیمه خالص",
                gross_key="premium",
                net_key="net_premium",
            )

            # Sheet دوم: وصولی
            collection_sheet = workbook.create_sheet("وصولی")
            fill_detail_sheet(
                sheet=collection_sheet,
                details=report.get("collection_details", []),
                gross_title="وصولی ناخالص",
                net_title="وصولی خالص",
                gross_key="collection_amount",
                net_key="net_collection_amount",
            )

            try:
                workbook.calculation.fullCalcOnLoad = True
                workbook.calculation.forceFullCalc = True
                workbook.calculation.calcMode = "auto"
            except Exception:
                pass

            workbook.save(file_path)

            messagebox.showinfo(
                "گزارش",
                "فایل Excel با موفقیت ایجاد شد.",
                parent=parent,
            )

        except Exception as e:
            messagebox.showerror(
                "خطا",
                f"خطا در ایجاد فایل اکسل:\n{e}",
                parent=parent,
            )
